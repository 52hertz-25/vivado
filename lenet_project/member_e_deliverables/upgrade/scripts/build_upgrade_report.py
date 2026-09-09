#!/usr/bin/env python3
"""Generate Member E's upgrade report from validated repository artifacts."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


UPGRADE = Path(__file__).resolve().parents[1]
OUT = UPGRADE / "report" / "Member_E_Joint_DSE_Report.docx"
JOINT = UPGRADE / "joint_dse"
MEMORY = UPGRADE / "memory_cache_candidate"
AXI = UPGRADE / "zynq_axi_integration"

BLUE = "17365D"
TEAL = "087E8B"
LIGHT = "D9EAF7"
PALE = "EEF5F9"
ORANGE = "E07A1F"
GRAY = "5B6573"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color="000000", size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True, "FFFFFF", font_size)
        set_cell_shading(table.rows[0].cells[i], BLUE)
        if widths:
            table.rows[0].cells[i].width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, "263238", font_size)
            if widths:
                cells[i].width = Cm(widths[i])
            if ridx % 2:
                set_cell_shading(cells[i], PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [("Title", 25, BLUE), ("Heading 1", 17, BLUE),
                               ("Heading 2", 12.5, TEAL), ("Heading 3", 10.5, GRAY)]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    for sec in doc.sections:
        header = sec.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = header.add_run("智能芯片设计 · LeNet 任务 2 · 成员 E 升级交付")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(GRAY)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = footer.add_run("成员 E（陈正堃）  |  ")
        rr.font.size = Pt(8)
        add_field(footer, "PAGE")


def add_badge(doc, label, text, fill):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.4)
    table.columns[1].width = Cm(12.8)
    set_cell_text(table.cell(0, 0), label, True, "FFFFFF", 10)
    set_cell_shading(table.cell(0, 0), fill)
    set_cell_text(table.cell(0, 1), text, False, "263238", 9.5)
    set_cell_shading(table.cell(0, 1), "F5F7F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def new_page(doc):
    doc.add_page_break()


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def build():
    validated = read_csv(JOINT / "results" / "joint_dse_validated.csv")
    configs = read_csv(JOINT / "joint_dse_configs.csv")
    memory = read_csv(MEMORY / "results" / "memory_cache_comparison.csv")

    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.title = "Member E Joint Bitwidth × PE/Pipeline DSE Upgrade Report"
    props.author = "陈正堃（成员 E）"
    props.subject = "智能芯片设计课程任务 2：LeNet FPGA 加速器"

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("成员 E 升级任务交付报告")
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("LeNet 联合位宽 × PE/Pipeline DSE、片上权重缓存与 Zynq AXI 集成")
    r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph()
    add_badge(doc, "课程任务", "任务 2：基于 HLS 的 LeNet FPGA 加速器（路线 B）", BLUE)
    add_badge(doc, "成员分工", "成员 E 陈正堃：HLS IP 封装、Vivado 综合/资源分析、自动搜索脚本", TEAL)
    add_badge(doc, "器件与约束", "XC7Z020-CLG400-1；统一目标周期 10.0 ns", ORANGE)
    add_badge(doc, "报告日期", "2026-09-09", GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("数据原则：仅报告仓库中可追溯的实测/既有工程结果；缺失项保持空值并标记阻塞。")

    # Executive summary
    new_page(doc)
    doc.add_heading("1. 交付结论", level=1)
    add_badge(doc, "主线 DSE", "建立 13 个联合配置；4 个 FP32 已在本机 Vitis HLS 2022.2 完成 CSim 与综合，9 个量化配置因成员 F 输入缺失而阻塞。", BLUE)
    add_badge(doc, "推荐配置", "D2（PE_MAC=5 + OUTPUT_COL pipeline）：243,101 cycles，10 ns 下 2.43101 ms，LUT 13,444，FF 16,887，DSP 7。", TEAL)
    add_badge(doc, "缓存候选", "两组均完成 Vitis HLS CSim/综合；缓存延迟增加 0.1037%、LUT +780、BRAM18K +10，结论为停止采用。", ORANGE)
    add_badge(doc, "AXI 集成", "实际运行 Vivado batch：强制校验 BD、重新导出实现报告并生成 XSA；26,721 条网络全部完成，WNS +0.415 ns。", GRAY)

    doc.add_heading("1.1 推荐依据", level=2)
    add_bullet(doc, "D2 相对 D0 的周期数减少 80.40%，是四个 FP32 候选中延迟最低者。")
    add_bullet(doc, "D1 虽做 pipeline，但 LUT/FF 增长较大且延迟仍高于 D2，被联合指标判为 dominated。")
    add_bullet(doc, "D3 增加输出通道并行后周期反而比 D2 增加 16.50%，说明当前瓶颈不适合继续盲目增加 OC PE。")
    add_bullet(doc, "D0 仍是最低资源候选；若资源优先于吞吐，它是保守备选。")

    doc.add_heading("1.2 不能越界声称的结果", level=2)
    p = doc.add_paragraph()
    p.add_run("INT8/INT6/INT4：").bold = True
    p.add_run("未收到成员 F 的量化源码、scale/zero-point 和真实准确率，故没有综合、准确率或 Pareto 数值。")
    p = doc.add_paragraph()
    p.add_run("缓存候选：").bold = True
    p.add_run("已完成 HLS 对照，实测没有性能收益，按停止规则保留为否定实验。")
    p = doc.add_paragraph()
    p.add_run("板级运行：").bold = True
    p.add_run("没有实机推理记录；AXI 部分是可复核的硬件集成包和软件调用模板。")

    # Background & matrix
    new_page(doc)
    doc.add_heading("2. 背景、分工与实验设计", level=1)
    doc.add_heading("2.1 宏观任务", level=2)
    doc.add_paragraph("课程任务书要求完成 LeNet FPGA 推理加速器：基础目标包括 MNIST 准确率、XC7Z020 综合、定点化依据、体系结构与存储说明；加分项要求至少三种额外位宽的真实准确率/资源/延迟结果、曲线及 Pareto 分析。本报告仅覆盖成员 E 的 Vivado 与 DSE 职责。")
    doc.add_heading("2.2 控制变量", level=2)
    add_table(doc, ["变量", "统一设置", "作用"], [
        ["器件", "xc7z020clg400-1", "与课程板级目标一致"],
        ["目标时钟", "10.0 ns (100 MHz)", "避免跨配置时钟不一致"],
        ["算子/数据", "Conv2 + 同一 testbench/model", "隔离 PE/Pipeline 影响"],
        ["精度维度", "FP32 / INT8 / INT6 / INT4", "量化精度—硬件代价联合搜索"],
        ["结构维度", "D0 / D1 / D2；D3 负对照", "观察 pipeline 与 PE 扩展"],
    ], [3.0, 5.0, 8.0], 8.7)
    doc.add_heading("2.3 13 个配置的执行状态", level=2)
    rows = []
    for c in configs:
        rows.append([c["config_id"], c["precision"], c["pe_variant"], c["pipeline_mode"],
                     "可验证" if c["planned_status"] == "READY" else "阻塞：缺 F 输入"])
    add_table(doc, ["配置 ID", "位宽", "结构", "Pipeline/PE", "状态"], rows,
              [5.0, 1.5, 1.3, 5.0, 3.3], 7.2)

    # Measured result table
    new_page(doc)
    doc.add_heading("3. FP32 × PE/Pipeline 实测结果", level=1)
    doc.add_paragraph("以下四行由成员 E 在本机 Vitis HLS 2022.2 对成员 D 提供的 D0-D3 源码逐一执行 CSim 与 C Synthesis 后，从原始 XML 重新解析。`execution_status=VALIDATED`、`csim=PASS` 表示本次复跑成功。")
    rows = []
    for r in validated:
        if r["precision"] != "FP32":
            continue
        rows.append([r["pe_variant"], r["pipeline_mode"], f'{int(r["latency_cycles"]):,}',
                     f'{float(r["latency_ms_at_target"]):.5f}', f'{int(r["lut"]):,}',
                     f'{int(r["ff"]):,}', r["bram_18k"], r["dsp"], r["pareto_status"],
                     r["selection_role"] or "—"])
    add_table(doc, ["结构", "策略", "Cycles", "ms@10ns", "LUT", "FF", "BRAM", "DSP", "Pareto", "角色"],
              rows, [1.0, 4.2, 2.3, 1.7, 1.5, 1.6, 1.2, 1.0, 1.4, 2.0], 6.7)
    doc.add_heading("3.1 定量观察", level=2)
    add_bullet(doc, "D0 → D1：延迟从 1,240,076 降到 254,749 cycles（-79.46%），代价是 LUT +93.39%、FF +416.93%。")
    add_bullet(doc, "D1 → D2：延迟再降 4.57%，同时 LUT -16.96%、FF -37.77%；因此 D2 在延迟和资源上同时支配 D1。")
    add_bullet(doc, "D2 → D3：延迟上升 16.50%，LUT 上升 15.11%，只有 FF 降低；D3 可作为结构负对照。")
    add_bullet(doc, "D0 的 estimated clock 8.280 ns；D1/D2/D3 约 10.20 ns。10 ns 下的 ms 列用于统一周期换算，不等价于每个 HLS 配置均满足估计时钟。")
    doc.add_heading("3.2 准确率字段的正确解释", level=2)
    doc.add_paragraph("CSV 中 99.30% 是成员 B 的 FP32 软件基线引用，并结合成员 D 的 CSim PASS 记录用于功能一致性说明；它不是这四个 Conv2 HLS 结构分别跑完整 MNIST 测试集得到的独立准确率。量化精度没有任何填充值。")

    # Figure
    new_page(doc)
    doc.add_heading("4. 联合 Pareto 结果", level=1)
    chart = JOINT / "figures" / "joint_pareto_accuracy_latency_resource.png"
    doc.add_picture(str(chart), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph("图 1  FP32 候选的延迟—资源 Pareto；量化配置因输入缺失不进入曲线。")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    doc.add_heading("4.1 选择规则", level=2)
    doc.add_paragraph("资源压力取 LUT、FF、BRAM、DSP 占器件容量百分比的最大值；Pareto 要求不存在另一个配置在延迟与资源压力上均不差且至少一项更优。平衡评分对延迟与资源压力分别做 min-max 归一化后等权相加。D2 的平衡评分 0.303153 最低，因此被标记为 `RECOMMENDED_BALANCED`。")
    doc.add_heading("4.2 量化搜索的待补闭环", level=2)
    doc.add_paragraph("成员 F 提交每个精度的定点 C/C++ 类型、scale/zero-point、量化权重/偏置和真实 MNIST 准确率后，可直接补入 9 个已固定 ID 的配置并运行同一脚本。这样配置命名、时钟、器件、输出字段与 Pareto 判定均无需重新设计。")

    # Memory cache
    new_page(doc)
    doc.add_heading("5. 片上权重缓存候选", level=1)
    doc.add_heading("5.1 改动", level=2)
    doc.add_paragraph("以 D2 为基线，新建 `lenet_conv_cache.cpp`：函数进入后将 2,400 个 Conv2 权重从 `m_axi` 预取到静态本地数组 `weight_cache`，以双端口 BRAM 绑定并做 cyclic factor=5 分割，再交给 PE_MAC=5 的卷积核。原 D2 源码保持不变。")
    add_table(doc, ["对象", "关键 pragma/结构", "预期作用"], [
        ["weight_cache", "BIND_STORAGE ram_2p bram", "将重复权重访问转为片上存储"],
        ["weight_cache", "ARRAY_PARTITION cyclic factor=5", "匹配 5 路 MAC 并行读"],
        ["CACHE_CONV2_WEIGHTS", "PIPELINE II=1", "连续搬运 2,400 个权重"],
        ["卷积主体", "保持 D2 PE_MAC=5", "控制变量，只测试缓存策略"],
    ], [3.0, 6.5, 6.5], 8.5)
    doc.add_heading("5.2 数值一致性", level=2)
    rows = [[r["config_id"], r["status"], r["csim"], r["max_abs_error"], r["decision"]] for r in memory]
    add_table(doc, ["配置", "状态", "测试", "最大绝对误差", "决策"], rows,
              [4.2, 4.2, 2.0, 2.4, 3.6], 7.5)
    doc.add_paragraph("两个版本均完成 Vitis HLS CSim 与 C Synthesis，CSim 最大绝对误差 2.7756691e-3，小于 5e-3 阈值。缓存版本延迟由 243,101 增至 243,353 cycles（+0.1037%），LUT 增加 780、FF 减少 4,766、BRAM18K 增加 10、DSP 不变。因此决策为 `STOP_NO_PERFORMANCE_GAIN`，不替换主线 D2。")
    doc.add_heading("5.3 后续唯一有效判据", level=2)
    add_bullet(doc, "先运行 baseline 与 cache 两个 solution 的 CSim；任一失败立即停止综合。")
    add_bullet(doc, "两者使用相同器件、10 ns 时钟、源码模型和 testbench。")
    add_bullet(doc, "只有当 cache 延迟/II 改善且资源增量可接受时，才进入主线；否则保留为否定实验。")

    # AXI
    new_page(doc)
    doc.add_heading("6. Zynq AXI 集成交付", level=1)
    fig = AXI / "figures" / "block_design_connection_map.png"
    doc.add_picture(str(fig), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph("图 2  从 `lenet_system.bd` 重建的连接图（不是 GUI 截图）。")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    doc.add_heading("6.1 地址与控制", level=2)
    add_table(doc, ["地址/偏移", "寄存器", "用途"], [
        ["0x4000_0000", "AXI-Lite control base", "PS M_AXI_GP0 访问 HLS 控制口"],
        ["+0x00", "AP_CTRL", "start/done/idle/ready/auto-restart"],
        ["+0x10/+0x14", "feature_in (64-bit)", "输入 DDR 物理地址"],
        ["+0x1C/+0x20", "weight (64-bit)", "权重 DDR 物理地址"],
        ["+0x28/+0x2C", "bias (64-bit)", "偏置 DDR 物理地址"],
        ["+0x34/+0x38", "feature_out (64-bit)", "输出 DDR 物理地址"],
    ], [3.3, 5.2, 7.4], 8.0)

    # AXI evidence
    new_page(doc)
    doc.add_heading("7. Vivado 实现证据", level=1)
    doc.add_paragraph("本节来自 2026-09-09 实际执行的 Vivado 2022.2 batch：打开工程，对 `lenet_system.bd` 执行 `validate_bd_design -force`，加载 `impl_1`，重新生成 route/timing/hierarchical utilization/power 报告，并成功导出 XSA。")
    add_table(doc, ["指标", "结果", "判定/含义"], [
        ["时钟", "10.000 ns / 100 MHz", "与联合 DSE 控制变量一致"],
        ["WNS / TNS", "+0.415 ns / 0.000 ns", "setup timing PASS"],
        ["WHS / THS", "+0.028 ns / 0.000 ns", "hold timing PASS"],
        ["可路由网络", "26,721 / 26,721", "全部完成"],
        ["路由错误", "0", "PASS"],
        ["LUT / FF", "11,612 / 17,612", "整套 Zynq 系统 placed utilization"],
        ["BRAM / DSP", "37 / 5", "整套 Zynq 系统 placed utilization"],
    ], [4.0, 5.1, 7.0], 8.7)
    doc.add_heading("7.1 IP 固定与可追溯性", level=2)
    add_bullet(doc, "IP VLNV：`user.org:hls:lenet_full_top:1.0`；归档 component.xml、export.zip、IP zip 和 HLS 驱动。")
    add_bullet(doc, "BD 源、地址表、PS 轮询伪代码、实现报告和历史 BD 校验日志一并归档。")
    add_bullet(doc, "`artifact_manifest_sha256.csv` 对包内文件给出字节数和 SHA-256，可检查交接后是否变化。")
    add_bullet(doc, "`rerun_vivado_reports.tcl` 可重新 validate BD，并从 routed DCP 导出 route/timing/utilization/power/XSA。")

    # Reproducibility
    new_page(doc)
    doc.add_heading("8. 复现入口、文件对应与责任边界", level=1)
    add_table(doc, ["工作项", "核心文件", "用途"], [
        ["联合矩阵", "joint_dse/joint_dse_configs.csv", "13 个稳定配置 ID 与依赖状态"],
        ["HLS 编排", "joint_dse/scripts/run_joint_dse.tcl", "逐配置 CSim→综合，失败即停"],
        ["结果收集", "joint_dse/scripts/collect_joint_results.py", "解析本次 XML、Pareto 和图"],
        ["验证数据", "joint_dse/results/joint_dse_validated.csv", "唯一结构化结果表"],
        ["缓存源码", "memory_cache_candidate/src/lenet_conv_cache.cpp", "D2 + 本地分割权重缓存"],
        ["缓存对照", "memory_cache_candidate/scripts/run_memory_cache.tcl", "baseline/cache 受控实验"],
        ["AXI 证据", "zynq_axi_integration/", "IP、BD、地址表、驱动、实现报告"],
        ["本报告", "report/Member_E_Joint_DSE_Report.docx", "成员 E 升级任务总交付"],
    ], [3.0, 7.0, 6.0], 8.0)
    doc.add_heading("8.1 推荐执行命令", level=2)
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    code = (
        "# 重新收集联合结果与生成图\n"
        "E:\\anaconda\\python.exe joint_dse\\scripts\\collect_joint_results.py\n\n"
        "# 缓存候选 HLS（Vitis HLS Tcl Console）\n"
        "vitis_hls -f memory_cache_candidate\\scripts\\run_memory_cache.tcl\n\n"
        "# Zynq 报告（Vivado Tcl）\n"
        "vivado -mode batch -source zynq_axi_integration\\scripts\\rerun_vivado_reports.tcl"
    )
    r = p.add_run(code)
    r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    r.font.size = Pt(8.5)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    doc.add_heading("8.2 当前阻塞项", level=2)
    add_bullet(doc, "成员 F：INT8/INT6/INT4 量化源码、scale/zero-point、量化参数文件和真实 MNIST 准确率。")
    add_bullet(doc, "路径兼容：Vitis HLS 2022.2 需在纯英文临时路径运行；本次已按该方式完成并归档原始报告。")
    add_bullet(doc, "成员 C：板级 PS 工程、DDR 缓冲区与缓存维护 API、device ID、实机测试结果。")
    doc.add_heading("8.3 最终结论", level=2)
    doc.add_paragraph("在现有输入范围内，成员 E 已实际完成 4 组 FP32 HLS CSim/综合、联合 Pareto、缓存基线/候选 HLS 对照、Vivado BD 校验与实现报告/XSA 导出，以及 Zynq AXI 交接包。主线推荐 D2，资源保守备选 D0；缓存候选因无性能收益停止。9 个量化配置继续等待成员 F 的真实输入，不以估计数补齐。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
