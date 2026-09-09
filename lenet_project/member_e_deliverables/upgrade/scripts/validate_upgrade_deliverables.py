#!/usr/bin/env python3
"""Fail-fast validation for Member E upgrade deliverables."""

from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def require(condition, message: str):
    if not condition:
        raise AssertionError(message)


def main():
    configs = rows(ROOT / "joint_dse" / "joint_dse_configs.csv")
    results = rows(ROOT / "joint_dse" / "results" / "joint_dse_validated.csv")
    require(len(configs) == 13, "joint config matrix must contain 13 rows")
    require(len({row["config_id"] for row in configs}) == 13, "config IDs must be unique")

    measured = [row for row in results if row["execution_status"] == "VALIDATED"]
    blocked = [row for row in results if row["execution_status"].startswith("BLOCKED_")]
    require(len(measured) == 4 and len(blocked) == 9,
            "expected 4 fresh HLS rows and 9 blocked quantized rows")
    require(all(row["csim"] == "PASS" for row in measured),
            "all four fresh FP32 CSim runs must pass")
    numeric_fields = ["accuracy_percent", "latency_cycles", "lut", "ff", "bram_18k", "dsp"]
    require(all(not row[field] for row in blocked for field in numeric_fields),
            "blocked quantized rows must not contain invented numeric results")
    chosen = [row for row in results if row["selection_role"] == "RECOMMENDED_BALANCED"]
    require(len(chosen) == 1 and chosen[0]["pe_variant"] == "D2",
            "D2 must be the balanced recommendation")

    joint_log = (ROOT / "joint_dse" / "evidence" / "fresh_vitis_hls_run_20260909.txt").read_text(
        encoding="utf-8", errors="replace")
    for config_id in [
        "fp32_d0_no_pipeline_10ns",
        "fp32_d1_pipeline_10ns",
        "fp32_d2_pe5_pipeline_10ns",
        "fp32_d3_pe5_oc2_pipeline_10ns",
    ]:
        require(f"MEMBER_E_CSIM_PASS {config_id}" in joint_log,
                f"missing CSim success marker: {config_id}")
        require(f"MEMBER_E_CSYNTH_PASS {config_id}" in joint_log,
                f"missing synthesis success marker: {config_id}")

    memory = rows(ROOT / "memory_cache_candidate" / "results" / "memory_cache_comparison.csv")
    require(len(memory) == 2, "memory experiment must have baseline and cache rows")
    require(all(row["status"] == "VALIDATED" for row in memory),
            "both memory variants must come from fresh HLS runs")
    require(all(row["csim"] == "PASS" for row in memory),
            "both memory CSim runs must pass")
    memory_by_id = {row["config_id"]: row for row in memory}
    baseline = memory_by_id["d2_external_weight_baseline"]
    cache = memory_by_id["d2_local_weight_cache"]
    require(baseline["decision"] == "CONTROL", "baseline must remain the control")
    require(cache["decision"] == "STOP_NO_PERFORMANCE_GAIN",
            "cache candidate must stop because latency did not improve")
    require(int(baseline["latency_cycles"]) == 243101, "unexpected baseline latency")
    require(int(cache["latency_cycles"]) == 243353, "unexpected cache latency")
    require(int(cache["bram_18k"]) == 10, "cache candidate must infer 10 BRAM18K")

    memory_log = (ROOT / "memory_cache_candidate" / "evidence" / "fresh_vitis_hls_run_20260909.txt").read_text(
        encoding="utf-8", errors="replace")
    for config_id in ["d2_external_weight_baseline", "d2_local_weight_cache"]:
        require(f"MEMBER_E_CACHE_CSIM_PASS {config_id}" in memory_log,
                f"missing memory CSim marker: {config_id}")
        require(f"MEMBER_E_CACHE_CSYNTH_PASS {config_id}" in memory_log,
                f"missing memory synthesis marker: {config_id}")

    axi = ROOT / "zynq_axi_integration"
    for relative in [
        "README.md",
        "ip/lenet_full_top_v1_0/component.xml",
        "ip/lenet_full_top_v1_0/export.zip",
        "bd/lenet_system.bd",
        "docs/address_and_register_map.md",
        "docs/ps_call_flow.md",
        "figures/block_design_connection_map.png",
        "evidence/bd_validation_history.txt",
        "evidence/fresh_vivado_rerun_20260909.txt",
        "reports/implementation_summary.csv",
        "reports/rerun/route_status.rpt",
        "reports/rerun/timing_summary.rpt",
        "reports/rerun/utilization_hierarchical.rpt",
        "reports/rerun/power.rpt",
        "reports/rerun/lenet_system_wrapper.xsa",
        "scripts/rerun_vivado_reports.tcl",
        "scripts/run_vivado_reports.ps1",
    ]:
        require((axi / relative).is_file(), f"missing AXI artifact: {relative}")

    vivado_log = (axi / "evidence" / "fresh_vivado_rerun_20260909.txt").read_text(
        encoding="utf-8", errors="replace")
    require("MEMBER_E_VIVADO_RERUN_PASS" in vivado_log,
            "missing Vivado success marker")
    route_text = (axi / "reports" / "rerun" / "route_status.rpt").read_text(
        encoding="utf-8", errors="replace")
    timing_text = (axi / "reports" / "rerun" / "timing_summary.rpt").read_text(
        encoding="utf-8", errors="replace")
    require("# of fully routed nets............. :       26721 :" in route_text
            and "# of nets with routing errors.......... :           0 :" in route_text,
            "fresh route report does not prove full routing")
    require("0.415" in timing_text and "0.028" in timing_text,
            "fresh timing report is missing expected slack values")
    for archive in [
        axi / "ip" / "lenet_full_top_v1_0" / "export.zip",
        axi / "reports" / "rerun" / "lenet_system_wrapper.xsa",
    ]:
        with zipfile.ZipFile(archive) as package:
            require(package.testzip() is None, f"corrupt archive: {archive}")

    manifest = rows(axi / "artifact_manifest_sha256.csv")
    for row in manifest:
        path = axi / row["relative_path"]
        require(path.is_file(), f"manifest path missing: {path}")
        require(int(row["size_bytes"]) == path.stat().st_size, f"size mismatch: {path}")
        require(hashlib.sha256(path.read_bytes()).hexdigest() == row["sha256"],
                f"hash mismatch: {path}")

    docx = ROOT / "report" / "Member_E_Joint_DSE_Report.docx"
    pdf = ROOT / "report" / "Member_E_Joint_DSE_Report.pdf"
    require(docx.stat().st_size > 100_000, "DOCX unexpectedly small")
    require(pdf.stat().st_size > 100_000, "PDF unexpectedly small")
    with zipfile.ZipFile(docx) as package:
        require(package.testzip() is None, "DOCX ZIP package is corrupt")
        require("word/document.xml" in package.namelist(), "DOCX missing main document part")
    document = Document(docx)
    full_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    for required in ["交付结论", "联合 Pareto 结果", "片上权重缓存候选", "Zynq AXI 集成交付", "Vivado 实现证据"]:
        require(required in full_text, f"DOCX missing section: {required}")
    require(len(document.tables) >= 8, "DOCX should contain at least 8 tables")
    require(len(document.inline_shapes) == 2, "DOCX should contain exactly 2 figures")
    require("2.43101" in full_text and "0.415" in full_text,
            "DOCX missing key fresh measurements")

    summary = {
        "status": "PASS",
        "joint_configs": 13,
        "fresh_hls_fp32": 4,
        "blocked_quantized": 9,
        "recommended": chosen[0]["config_id"],
        "memory_hls_validated": 2,
        "memory_decision": cache["decision"],
        "vivado_rerun": "PASS",
        "xsa_exported": True,
        "docx_tables": len(document.tables),
        "docx_figures": len(document.inline_shapes),
        "axi_manifest_entries": len(manifest),
    }
    output = ROOT / "validation_summary.json"
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
