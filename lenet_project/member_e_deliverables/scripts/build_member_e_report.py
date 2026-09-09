from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIG = ROOT / "figures"
OUT = ROOT / "report" / "成员E_LeNet_DSE实验报告.docx"

NAVY = "14213D"
BLUE = "2E6797"
GOLD = "C99116"
LIGHT_BLUE = "E8F0F8"
LIGHT_GOLD = "FBF2DC"
LIGHT_GRAY = "F2F5F8"
MID_GRAY = "AAB7C5"
TEXT = RGBColor(20, 33, 61)


def read_csv(name: str) -> list[dict[str, str]]:
    with (DATA / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = NAVY, size: float = 8.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        set_cell_fill(table.rows[0].cells[idx], NAVY)
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="FFFFFF", size=8)
        if widths:
            table.rows[0].cells[idx].width = Inches(widths[idx])
    set_repeat_table_header(table.rows[0])
    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, value in enumerate(values):
            set_cell_fill(cells[c_idx], "FFFFFF" if r_idx % 2 == 0 else LIGHT_GRAY)
            set_cell_text(cells[c_idx], value, bold=(r_idx == 0 and values[0].startswith("10 ns")), size=7.8)
            if widths:
                cells[c_idx].width = Inches(widths[c_idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_kicker(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.name = "Aptos"


def add_title(doc: Document, text: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = TEXT
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if subtitle:
        p2 = doc.add_paragraph(subtitle)
        p2.style = doc.styles["Subtitle"]


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        first, rest = text[: len(bold_lead)], text[len(bold_lead) :]
        p.add_run(first).bold = True
        p.add_run(rest)
    else:
        p.add_run(text)


def add_bullets(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(line)


def add_callout(doc: Document, title: str, body: str, color: str = LIGHT_GOLD) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, color)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_picture(doc: Document, name: str, width: float = 6.6, caption: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / name), width=Inches(width))
    if caption:
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.style = doc.styles["Caption"]


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    for name, size, color in [("Title", 30, NAVY), ("Subtitle", 12, "55657A"), ("Heading 1", 19, NAVY), ("Heading 2", 14, BLUE)]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
    styles["Caption"].font.name = "Microsoft YaHei"
    styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.color.rgb = RGBColor.from_string("64748B")
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "LE NET · MEMBER E · DSE & VIVADO"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string("708096")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("成员 E 实验交付  ·  ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)


clock_rows = read_csv("clock_dse_validated.csv")
struct_rows = read_csv("structural_dse_validated.csv")
doc = Document()
configure_document(doc)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(42)
add_kicker(doc, "智能芯片设计 · 任务 2 · LeNet")
cover = doc.add_paragraph()
cover.paragraph_format.space_after = Pt(12)
run = cover.add_run("LeNet HLS 与 Vivado\n设计空间探索报告")
run.bold = True
run.font.size = Pt(30)
run.font.color.rgb = TEXT
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
sub = doc.add_paragraph("成员 E：陈正堃  ·  Vitis HLS / Vivado 2022.2  ·  XC7Z020")
sub.style = doc.styles["Subtitle"]
doc.add_paragraph().paragraph_format.space_after = Pt(48)
add_callout(doc, "最终选择：clk_10ns_baseline", "该配置在 5 组时钟约束中同时具有最少调度周期、最低按目标周期换算的总延迟和最低 LUT；Vivado post-route 时序通过。")
doc.add_paragraph().paragraph_format.space_after = Pt(82)
meta = doc.add_paragraph("交付内容：两张核验 CSV · 三张 DSE 图 · 原始报告 · Vivado 证据 · 可复现实验包 · 3 页答辩 PPT")
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.runs[0].font.size = Pt(9)
meta.runs[0].font.color.rgb = RGBColor.from_string("64748B")

doc.add_page_break()
add_kicker(doc, "Executive summary")
add_title(doc, "结论先行", "成员 E 只负责 HLS/Vivado 设计空间探索与实现证据归档")
add_body(doc, "本次工作复核了 5 组时钟约束和 3 组 MAC 循环展开实验，数据均由原始 Vitis HLS `csynth.xml` 重新提取，并与项目已有 CSV 逐项比对。所有比较均使用同一 LeNet FP32 网络、同一 XC7Z020 器件和相同接口边界。")
add_bullets(doc, [
    "时钟约束从 10 ns 收紧到 7 ns 时，HLS 估计 Fmax 由 120.77 MHz 提升到 185.98 MHz，但调度周期从 1,647,854 增至 2,847,628。",
    "10 ns 基线按目标时钟换算的总延迟为 16.479 ms，低于其余 4 个配置的 19.115–20.374 ms，同时 LUT 最少（17,108）。",
    "baseline、Unroll ×2、Unroll ×5 的延迟和资源完全一致，说明当前简单展开没有改变综合后的有效并行度。",
    "Vivado post-route：WNS +0.415 ns、WHS +0.028 ns；LUT 11,612、寄存器 17,612、BRAM 37 tiles、DSP 5；总片上功耗估算 1.848 W。",
])
add_callout(doc, "停止条件已满足", "E任务优化方案规定：若现有 DSE 证据足以形成独立贡献，可停止继续改代码。因此本次不引入会影响其他成员接口或数值口径的新优化。", LIGHT_BLUE)

doc.add_page_break()
add_kicker(doc, "01 · scope & method")
add_title(doc, "实验范围与核验方法")
add_heading(doc, "成员 E 的交付边界", 2)
add_body(doc, "负责 HLS 时钟约束 DSE、结构探索、Pareto 解释、最终 Vivado 布线后结果归档与答辩材料；不承担量化训练、接口集成或其他成员的模块任务。")
add_heading(doc, "实验矩阵", 2)
add_table(doc, ["维度", "配置", "唯一变量", "判定"], [
    ["Clock DSE", "10 / 9 / 8 / 7.5 / 7 ns", "create_clock 目标周期", "延迟、Fmax、LUT/FF/BRAM/DSP"],
    ["Structural DSE", "baseline / Unroll ×2 / ×5", "MAC 循环展开因子", "是否产生有效并行度"],
    ["Vivado", "10 ns baseline", "最终候选", "post-route 时序、资源、功耗"],
], [1.0, 2.0, 1.8, 2.1])
add_heading(doc, "证据链", 2)
add_bullets(doc, [
    "HLS：每个配置的 `solution1/syn/report/lenet_full_top_csynth.xml`。",
    "Vivado：timing summary、placed utilization、routed power、route status 原始报告。",
    "自动核验：`build_member_e_results.py` 重新提取所有指标并检查与原汇总 CSV 一致；不一致即终止。",
    "图表：从核验后的 CSV 生成，基线用金色高亮，每个点均直接标注。",
])
add_callout(doc, "单位约定", "HLS 使用 BRAM_18K；Vivado 使用 Block RAM Tile。两者是不同报告口径，不能直接做数值差。")

doc.add_page_break()
add_kicker(doc, "02 · clock DSE")
add_title(doc, "收紧时钟并未缩短端到端延迟")
table_rows = []
for row in clock_rows:
    cfg = {"clk_10ns_baseline": "10 ns 基线", "clk_9ns": "9 ns", "clk_8ns": "8 ns", "clk_7_5ns": "7.5 ns", "clk_7ns": "7 ns"}[row["Config"]]
    table_rows.append([
        cfg, f'{float(row["Estimated_ns"]):.3f}', f'{float(row["Estimated_Fmax_MHz"]):.2f}',
        f'{int(row["Latency_cycles"]):,}', f'{float(row["Latency_ms_at_target"]):.3f}',
        f'{int(row["LUT"]):,}', f'{int(row["FF"]):,}', f'{int(row["BRAM_18K"]):d} / {int(row["DSP"]):d}',
    ])
add_table(doc, ["配置", "估计周期\nns", "估计 Fmax\nMHz", "延迟\ncycles", "目标延迟\nms", "LUT", "FF", "BRAM18K\n/ DSP"], table_rows,
          [0.83, 0.65, 0.72, 0.92, 0.75, 0.65, 0.65, 0.78])
add_picture(doc, "clock_latency_cycles.png", 6.55, "图 1  Target_ns 与最坏情况延迟周期；10 ns 基线调度周期最少")

doc.add_page_break()
add_kicker(doc, "02 · clock DSE")
add_title(doc, "10 ns 位于延迟—LUT 左下角")
add_picture(doc, "pareto_latency_lut.png", 6.55, "图 2  延迟周期与 LUT 的 Pareto 对比；左下方向更优")
add_body(doc, "10 ns 基线同时具有最低延迟周期（1,647,854）与最低 LUT（17,108），因此对其余 4 个时钟配置形成弱支配。即使按各自目标周期换算，10 ns 的 16.479 ms 仍低于其他配置。")
add_heading(doc, "机制解释", 2)
add_body(doc, "时钟约束变紧会推动工具提高组合路径速度，所以估计 Fmax 上升；但为了满足更短周期，HLS 可能插入更多流水级、延长浮点运算调度或改变资源复用时序，使总周期数明显增加。频率收益不足以抵消周期数增加，端到端延迟反而恶化。")
add_callout(doc, "选择依据", "保留 clk_10ns_baseline：它不是频率最高的点，但在当前架构和接口下提供最低总延迟、最低 LUT，并已有完整 Vivado post-route 证据。")

doc.add_page_break()
add_kicker(doc, "03 · structural DSE")
add_title(doc, "单独展开 MAC 循环没有产生收益")
struct_table = []
for row in struct_rows:
    label = {"struct_baseline": "baseline", "mac_unroll_2": "Unroll ×2", "mac_unroll_5": "Unroll ×5"}[row["Config"]]
    struct_table.append([label, row["CSim"], f'{int(row["Latency_cycles"]):,}', f'{int(row["LUT"]):,}', f'{int(row["FF"]):,}', f'{int(row["BRAM_18K"]):d}', f'{int(row["DSP"]):d}'])
add_table(doc, ["配置", "CSim", "Latency cycles", "LUT", "FF", "BRAM18K", "DSP"], struct_table,
          [1.0, 0.65, 1.15, 0.75, 0.75, 0.75, 0.55])
add_picture(doc, "structural_dse_comparison.png", 6.65, "图 3  三组结构实验的延迟、LUT 与 DSP 完全相同")
add_heading(doc, "为什么结果相同（基于报告与代码结构的推断）", 2)
add_bullets(doc, [
    "四个 `m_axi` 外存接口带来的访问等待可能限制数据供给；MAC 循环展开无法绕过存储带宽瓶颈。",
    "FP32 运算单元仍被综合器共享，展开指令未必实例化额外乘加资源；DSP 数固定为 5 支持这一判断。",
    "上层循环、函数调用和依赖关系决定总体调度，局部 MAC 循环不是主导临界路径。",
])

doc.add_page_break()
add_kicker(doc, "04 · Vivado post-route")
add_title(doc, "最终基线已完成布线后闭环")
add_table(doc, ["类别", "指标", "结果", "判定"], [
    ["时序", "WNS / WHS", "+0.415 / +0.028 ns", "满足 setup/hold"],
    ["资源", "LUT / Registers", "11,612 / 17,612", "21.83% / 16.55%"],
    ["资源", "BRAM tiles / DSP", "37 / 5", "26.43% / 2.27%"],
    ["功耗", "Total On-Chip", "1.848 W", "Medium confidence"],
], [0.8, 1.35, 1.45, 2.25])
add_picture(doc, "vivado_timing_report_excerpt.png", 6.55, "图 4  Vivado 2022.2 timing summary 原始报告摘录")

doc.add_page_break()
add_kicker(doc, "04 · Vivado post-route")
add_title(doc, "资源与功耗证据")
add_picture(doc, "vivado_resource_report_excerpt.png", 6.4, "图 5  Vivado placed utilization 原始报告摘录")
add_picture(doc, "vivado_power_report_excerpt.png", 6.4, "图 6  Vivado routed power 原始报告摘录")
add_callout(doc, "功耗口径限制", "1.848 W 未使用仿真活动文件，报告 Confidence Level 为 Medium，其中 PS7 约 1.530 W。因此只能写作 Vivado 估算，不能写成板级实测功耗。", LIGHT_BLUE)

doc.add_page_break()
add_kicker(doc, "05 · decision & next step")
add_title(doc, "交付结论与后续条件")
add_heading(doc, "最终结论", 2)
add_bullets(doc, [
    "选择 10 ns 基线作为成员 E 的最终方案；不继续以更紧的时钟约束换取名义 Fmax。",
    "单独应用 MAC Unroll ×2/×5 无效，不应把它描述为性能优化；该负结果本身说明了架构瓶颈。",
    "最终 Vivado 实现 fully routed 且时序通过，资源与功耗均有可追溯原始报告。",
])
add_heading(doc, "仅在教师明确要求“必须改代码”时", 2)
add_body(doc, "新建独立候选版本，不覆盖基线；把权重缓存到片上数组，对缓存维度使用 ARRAY_PARTITION，并设计明确的并行 PE/MAC 数据通路。该候选必须依次通过 CSim 数值一致性、HLS 延迟/资源比较、导出 IP 和 Vivado post-route，才可称为有效优化。")
add_heading(doc, "可复现文件", 2)
add_bullets(doc, [
    "`data/clock_dse_validated.csv`、`data/structural_dse_validated.csv`、`data/vivado_post_route_summary.csv`：最终汇总。",
    "`data/validation_report.md`：逐项核验记录。",
    "`scripts/build_member_e_results.py`：提取报告、校验 CSV、生成图表。",
    "`scripts/build_vivado_evidence.mjs`：从原始 Vivado 报告生成可读证据图。",
    "`dse_package/`：配置、指令、源码、模型、变体、脚本与结果快照。",
    "`source_reports/`：10 ns HLS csynth 报告及 Vivado post-route 原始报告；复现命令见交付目录 README。",
])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
